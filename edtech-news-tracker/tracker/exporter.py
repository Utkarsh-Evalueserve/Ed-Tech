"""
Writes the final timestamped Excel workbook:
  - Sheet 'News': one row per (deduplicated) article.
  - Sheet 'Run Log': companies/geographies searched, articles found,
    zero-news companies, failed queries, and run duration.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import List

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .relevance import Article

NEWS_COLUMNS = [
    "Company", "Geography", "News Category", "Headline (English)",
    "English Summary", "Source Name", "Publication Date", "URL",
    "Relevance Score", "Review Flag",
]


@dataclass
class RunLog:
    start_date: str
    end_date: str
    companies_searched: List[str] = field(default_factory=list)
    geographies_searched: List[str] = field(default_factory=list)
    articles_found: int = 0
    companies_with_zero_news: List[str] = field(default_factory=list)
    failed_queries: List[str] = field(default_factory=list)
    duration_seconds: float = 0.0


def _style_header(ws) -> None:
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill


def export(articles: List[Article], run_log: RunLog, output_dir: Path) -> Path:
    wb = openpyxl.Workbook()

    ws_news = wb.active
    ws_news.title = "News"
    ws_news.append(NEWS_COLUMNS)
    for art in sorted(articles, key=lambda a: (a.company, a.geography, -a.relevance_score)):
        ws_news.append([
            art.company, art.geography, art.category, art.headline_en,
            art.summary_en, art.source_name, art.publication_date or "",
            art.url, art.relevance_score, art.review_flag,
        ])
    _style_header(ws_news)
    for idx, width in enumerate([25, 15, 30, 45, 55, 30, 15, 45, 12, 20], start=1):
        ws_news.column_dimensions[get_column_letter(idx)].width = width
    ws_news.freeze_panes = "A2"

    ws_log = wb.create_sheet("Run Log")
    ws_log.append(["Metric", "Value"])
    _style_header(ws_log)
    rows = [
        ("Date range", f"{run_log.start_date} to {run_log.end_date}"),
        ("Run duration (seconds)", round(run_log.duration_seconds, 1)),
        ("Companies searched", len(run_log.companies_searched)),
        ("Geographies searched", len(run_log.geographies_searched)),
        ("Total articles found (after dedup)", run_log.articles_found),
        ("Companies with zero news found", len(run_log.companies_with_zero_news)),
        ("Failed queries (after retries)", len(run_log.failed_queries)),
    ]
    for r in rows:
        ws_log.append(r)

    ws_log.append([])
    ws_log.append(["Companies with zero news found:"])
    for c in run_log.companies_with_zero_news:
        ws_log.append(["", c])

    ws_log.append([])
    ws_log.append(["Failed queries (after retries):"])
    for f in run_log.failed_queries:
        ws_log.append(["", f])

    ws_log.column_dimensions["A"].width = 35
    ws_log.column_dimensions["B"].width = 60

    output_dir.mkdir(parents=True, exist_ok=True)
    filename = f"EdTech_News_{run_log.start_date}_to_{run_log.end_date}.xlsx"
    out_path = output_dir / filename
    wb.save(out_path)
    return out_path
